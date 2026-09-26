"""Build the curated Konoha Shido 1st Edition reference JSON from its workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

WORKBOOK = Path("data/reference/konoha_shido_1st_edition_master.xlsx")
OUTPUT = Path("data/reference/konoha_shido_1st_edition_master.json")
REQUIRED_HEADERS = ("No.", "Card", "Rarity", "Variant")


def _text(value: Any) -> str:
    return " ".join(unicodedata.normalize("NFC", str(value or "")).split())


def _key_digest(values: tuple[str, ...]) -> str:
    payload = json.dumps(values, ensure_ascii=True, separators=(",", ":"))
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def build_reference(workbook_path: Path = WORKBOOK) -> dict[str, Any]:
    workbook_bytes = workbook_path.read_bytes()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if "Checklist" not in workbook.sheetnames:
        raise ValueError("CHECKLIST_SHEET_MISSING")
    sheet = workbook["Checklist"]
    rows = sheet.iter_rows(values_only=True)
    header_row_number = None
    header_values: tuple[Any, ...] | None = None
    buffered_rows = []
    for row_number, row in enumerate(rows, start=1):
        values = tuple(row)
        if header_row_number is None and all(name in values for name in REQUIRED_HEADERS):
            header_row_number = row_number
            header_values = values
            continue
        if header_row_number is not None:
            buffered_rows.append((row_number, values))
    if header_row_number is None or header_values is None:
        raise ValueError("CHECKLIST_HEADERS_MISSING")

    header_indexes = {name: header_values.index(name) for name in REQUIRED_HEADERS}
    entries = []
    for row_number, values in buffered_rows:
        record = {
            name: _text(values[index]) if index < len(values) else ""
            for name, index in header_indexes.items()
        }
        if not any(record.values()):
            continue
        if not all(record.values()):
            raise ValueError(f"INCOMPLETE_CHECKLIST_ROW:{row_number}")
        identity = tuple(record[name] for name in REQUIRED_HEADERS)
        entries.append(
            {
                "reference_key": "kref_" + _key_digest(identity)[:56],
                "workbook_row": row_number,
                "collector_number": record["No."],
                "card_name": record["Card"],
                "rarity_raw": record["Rarity"],
                "variant_raw": record["Variant"],
            }
        )

    keys = [entry["reference_key"] for entry in entries]
    if len(entries) != 396:
        raise ValueError(f"CHECKLIST_ENTRY_COUNT:{len(entries)}")
    if len(keys) != len(set(keys)):
        raise ValueError("DUPLICATE_CHECKLIST_IDENTITY")

    return {
        "schema_version": "konoha-shido-reference-v1",
        "source_file": workbook_path.name,
        "source_sha256": hashlib.sha256(workbook_bytes).hexdigest(),
        "worksheet": "Checklist",
        "header_row": header_row_number,
        "headers_used": list(REQUIRED_HEADERS),
        "entry_count": len(entries),
        "entries": entries,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=WORKBOOK)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    result = build_reference(args.input)
    encoded = json.dumps(result, ensure_ascii=False, sort_keys=True, indent=2) + "\n"
    if args.check:
        if not args.output.exists() or args.output.read_text(encoding="utf-8") != encoded:
            raise SystemExit("REFERENCE_ARTIFACT_OUT_OF_DATE")
    else:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(encoded, encoding="utf-8", newline="\n")
    print(f"reference_entries={result['entry_count']}")
    print(f"source_sha256={result['source_sha256']}")


if __name__ == "__main__":
    main()
